"""Đọc file CSDL .xlsx → tools/schema_raw.json (lược đồ thô) + test/data_real.json (dữ liệu thật).

Dùng:  python3 tools/extract_xlsx.py <file.xlsx> [test/data_real.json]
(test/data.json là fixture cố định cho test logic; data_real.json = dữ liệu thật để smoke test)
Lược đồ thô gồm cho mỗi bảng LSB_*: tên cột, kiểu (theo ghi chú tiêu đề "Kiểu: …"),
PK (tiêu đề màu nâu A7650F), FK (dropdown x14 trỏ tới LSB_xxx!$A), list (dropdown trỏ tới
sheet Lists), formula (dòng 2 có công thức). Đọc trực tiếp XML của sheet vì openpyxl
bỏ mất dropdown dạng x14.
"""
import json, re, sys, zipfile, datetime, warnings, os
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
warnings.filterwarnings('ignore')

SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', 'data', 'LSB_CRM_Bidding_Project_DB_v4.3.xlsx')
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

wb_f = openpyxl.load_workbook(SRC)                   # công thức + ghi chú + màu
wb_v = openpyxl.load_workbook(SRC, data_only=True)   # giá trị đã tính
z = zipfile.ZipFile(SRC)

# tên sheet → file xml
wbxml = z.read('xl/workbook.xml').decode()
rels = z.read('xl/_rels/workbook.xml.rels').decode()
rid2t = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
rid2t.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
sheet_xml = {n: 'xl/' + rid2t[r].lstrip('/').replace('xl/', '') for n, r in re.findall(r'<sheet name="([^"]+)" sheetId="\d+" r:id="(rId\d+)"', wbxml)}

lists_ws = wb_f['Lists']
list_by_col = {}
lists = {}
for c in range(1, lists_ws.max_column + 1):
    name = lists_ws.cell(1, c).value
    if not name: continue
    list_by_col[get_column_letter(c)] = name
    lists[name] = [v for v in (lists_ws.cell(r, c).value for r in range(2, lists_ws.max_row + 1)) if v not in (None, '')]

def cols_of(sqref):
    out = set()
    for part in sqref.split():
        m = re.match(r'([A-Z]+)\d+(?::([A-Z]+)\d+)?', part)
        a = column_index_from_string(m.group(1)); b = column_index_from_string(m.group(2) or m.group(1))
        out.update(range(a, b + 1))
    return out

tables = {}
data = {}
for name in wb_f.sheetnames:
    if not name.startswith('LSB_'): continue
    ws, wv = wb_f[name], wb_v[name]
    xml = z.read(sheet_xml[name]).decode()
    fk, lst = {}, {}
    for m in re.finditer(r'<x14:dataValidation [^>]*type="list".*?<xm:f>([^<]*)</xm:f>.*?<xm:sqref>([^<]*)</xm:sqref>', xml, re.S):
        f, sq = m.group(1), m.group(2)
        mm = re.match(r"'?(LSB_\w+)'?!\$A\$", f)
        ml = re.match(r"'?Lists'?!\$([A-Z]+)\$", f)
        for c in cols_of(sq):
            if mm: fk[c] = mm.group(1)
            elif ml: lst[c] = list_by_col.get(ml.group(1))
    cols = []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if not h: continue
        note = ws.cell(1, c).comment.text if ws.cell(1, c).comment else ''
        mt = re.search(r'Kiểu:\s*(\w+)', note)
        fill = (ws.cell(1, c).fill.fgColor.rgb or '')[-6:] if ws.cell(1, c).fill and ws.cell(1, c).fill.fgColor else ''
        f2 = ws.cell(2, c).value
        cols.append(dict(name=str(h).strip(), type=mt.group(1) if mt else 'string', fill=fill,
                         fk=fk.get(c, ''), list=lst.get(c, ''), formula=isinstance(f2, str) and f2.startswith('=')))
    tables[name] = cols
    grid = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            fv = ws.cell(r, c).value
            vv = wv.cell(r, c).value
            f = fv if isinstance(fv, str) and fv.startswith('=') else ''
            v = vv if f else fv
            if isinstance(v, (datetime.datetime, datetime.date)):
                if not isinstance(v, datetime.datetime): v = datetime.datetime(v.year, v.month, v.day)
                v = {'$d': v.strftime('%Y-%m-%dT%H:%M:%S')}
            row.append({'v': v, 'f': f})
        grid.append(row)
    data[name] = grid

json.dump({'source': os.path.basename(SRC), 'tables': tables, 'lists': lists},
          open(os.path.join(HERE, 'schema_raw.json'), 'w'), ensure_ascii=False, indent=1)
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'test', 'data_real.json')
json.dump(data, open(OUT, 'w'), ensure_ascii=False)
print('bảng:', len(tables), '· cột:', sum(len(v) for v in tables.values()))
for t, cs in tables.items():
    print(' ', t, len(cs), 'PK=' + ','.join(c['name'] for c in cs if c['fill'] == 'A7650F'),
          'FK=' + ','.join(c['name'] + '→' + c['fk'] for c in cs if c['fk']))
