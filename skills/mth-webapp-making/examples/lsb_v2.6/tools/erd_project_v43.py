"""Dựng lại sheet ERD_Dictionary cho file LSB_CRM_Bidding_Project_DB_v4.3_Project.xlsx (11 bảng).

Dùng: python3 tools/erd_project_v43.py <vào.xlsx> <ra.xlsx> [erd.mmd]

- Giữ nguyên mọi sheet khác (sửa trực tiếp XML → không mất dropdown x14).
- Từ điển: chỉ các bảng có trong file (theo thứ tự sheet), cột đối chiếu với tiêu đề thật của sheet.
- Quan hệ: chỉ quan hệ mà cả 2 bảng đều có trong file (+ mục "Tham chiếu ngoài" không FK).
- Ghi thêm sơ đồ Mermaid erDiagram (tùy chọn).
"""
import re, sys, zipfile
from xml.sax.saxutils import escape
import openpyxl, warnings
warnings.filterwarnings('ignore')

SRC, DST = sys.argv[1], sys.argv[2]
MMD = sys.argv[3] if len(sys.argv) > 3 else None
wb = openpyxl.load_workbook(SRC)
tables = [n for n in wb.sheetnames if n.startswith('LSB_')]
headers = {t: [str(c.value).strip() for c in wb[t][1] if c.value] for t in tables}
old = list(wb['ERD_Dictionary'].iter_rows(values_only=True))

# --- 1. từ điển cột (lấy dòng cũ theo (bảng, cột), sửa ghi chú v4.3) ---
dic = {}
for r in old[1:]:
    if r[0] in tables and r[1] and not str(r[1]).startswith('LSB_') and r[2] in headers.get(r[0], []):
        dic.setdefault((r[0], r[2]), list(r[:9]) + [None] * (9 - len(r[:9])))
FIX = {
 ('LSB_project', 'Quotation_LSB-Code'): (None, 'mã báo giá nội bộ thắng thầu (vd 25B561) = LSB_Quotation.Quotation_LSB_Code ở file Bidding — nhập tay, duy nhất', None),
 ('LSB_project', 'Contract_VND'): ('Nhập / Auto', '= Contract_USD × 26000 (công thức mẫu dòng 2; app tự tính khi để trống)', 'v4.3: thêm công thức'),
 ('LSB_shipment', 'ETA_plan'): ('Nhập / Auto', '= ETD_plan + 14 ngày (công thức mẫu dòng 2; app tự tính khi để trống)', 'v4.3: thêm công thức'),
 ('LSB_shipment', 'ETA_actual'): ('Nhập / Auto', '= ETD_actual + 14 ngày (công thức mẫu dòng 2; app tự tính khi để trống) — nên nhập ngày thật', 'v4.3: thêm công thức'),
}
dict_rows = []
missing = []
for t in tables:
    for c in headers[t]:
        row = dic.get((t, c))
        if not row:
            missing.append((t, c)); row = [t, '', c, '', '', '', 'Nhập', '', 'v4.3_Project: bổ sung từ điển']
        if (t, c) in FIX:
            g, h, i = FIX[(t, c)]
            if g: row[6] = g
            if h: row[7] = h
            if i: row[8] = (str(row[8]) + '; ' + i) if row[8] else i
        dict_rows.append(row)

# --- 2. quan hệ ---
rels, in_rel = [], False
for r in old:
    if r[0] and str(r[0]).startswith('Quan hệ'): in_rel = True; continue
    if r[0] and str(r[0]).startswith('Đã bỏ'): in_rel = False
    if in_rel and r[0] in tables and r[1] in tables:
        rels.append([r[0], r[1], r[2], r[3], r[4]])
ext = [['LSB_Quotation (file Bidding)', 'LSB_project', 'Quotation_LSB-Code → Quotation_LSB_Code', '1 : 0..1 (Won)', 'không FK — nhập tay mã nội bộ; app không kiểm tra vì bảng không có trong file']]

# --- 3. bảng không có trong file + đã bỏ so với v3.1 ---
BID = ['LSB_Lead_Qualifying', 'LSB_Quotation', 'LSB_Bidding_strategy', 'LSB_Bidding_action_plan', 'LSB_Result_Lesson']
absent = [[t, 'quản lý ở file …_Bidding / bản gộp v4.3', 'bỏ khỏi file Project (v4.3_Project) — kéo theo bỏ các quan hệ tới bảng này'] for t in BID]
dropped, in_drop = [], False
for r in old:
    if r[0] and str(r[0]).startswith('Đã bỏ'): in_drop = True; continue
    if in_drop and r[0] in tables: dropped.append([r[0], r[1], r[2]])

# --- 4. dựng XML ---
out, n = [], 0
def row(vals, style, span):
    global n
    n += 1
    cells = ''
    for i, v in enumerate(vals):
        col = 'ABCDEFGHI'[i]
        if v in (None, ''): cells += '<c r="%s%d" s="%s"/>' % (col, n, style)
        else: cells += '<c r="%s%d" s="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (col, n, style, escape(str(v)))
    out.append('<row r="%d" spans="1:%d" x14ac:dyDescent="0.3">%s</row>' % (n, span, cells))
row(['Bảng', 'Nhóm', 'Cột', 'Kiểu', 'Khóa', 'Tham chiếu', 'Nhập / Auto', 'Ghi chú', 'Thay đổi (v4.1 → v4.3)'], 39, 9)
for r in dict_rows: row(r, 25, 9)
last_dict = n
row([], 25, 1); out.pop(); n -= 1  # (không cần dòng trống thật)
n += 1
row(['Quan hệ (%d) — chỉ giữa các bảng có trong file' % len(rels)], 40, 5)
row(['Bảng cha', 'Bảng con', 'Khóa nối', 'Bản số', 'Ghi chú'], 39, 5)
for r in rels: row(r, 25, 5)
n += 1
row(['Tham chiếu ngoài (không khóa ngoại)'], 40, 5)
row(['Nguồn', 'Bảng', 'Cột', 'Bản số', 'Ghi chú'], 39, 5)
for r in ext: row(r, 25, 5)
n += 1
row(['Bảng không có trong file Project (v4.3_Project)'], 40, 3)
row(['Bảng', 'Vị trí', 'Ghi chú'], 39, 3)
for r in absent: row(r, 25, 3)
n += 1
row(['Đã bỏ so với v3.1 (các bảng có trong file)'], 40, 3)
row(['Bảng', 'Cột', 'Ghi chú'], 39, 3)
for r in dropped: row(r, 25, 3)

zin = zipfile.ZipFile(SRC)
files = {x: zin.read(x) for x in zin.namelist()}
wbxml = files['xl/workbook.xml'].decode()
rid = re.search(r'<sheet name="ERD_Dictionary" sheetId="\d+" r:id="(rId\d+)"', wbxml).group(1)
relsxml = files['xl/_rels/workbook.xml.rels'].decode()
target = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid, relsxml).group(1)
sn = 'xl/' + target.lstrip('/').replace('xl/', '')
s = files[sn].decode()
s = re.sub(r'<sheetData>.*</sheetData>', lambda m: '<sheetData>' + ''.join(out) + '</sheetData>', s, flags=re.S)
s = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1:I%d"/>' % n, s)
s = re.sub(r'<autoFilter ref="[^"]*"', '<autoFilter ref="A1:I%d"' % last_dict, s)
s = re.sub(r'<mergeCells.*?</mergeCells>', '', s, flags=re.S)
files[sn] = s.encode()
wbxml = re.sub(r'(ERD_Dictionary!\$A\$1:\$I\$)\d+', lambda m: m.group(1) + str(last_dict), wbxml)
files['xl/workbook.xml'] = wbxml.encode()
if 'xl/calcChain.xml' in files:
    del files['xl/calcChain.xml']
    files['[Content_Types].xml'] = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', '', files['[Content_Types].xml'].decode()).encode()
    files['xl/_rels/workbook.xml.rels'] = re.sub(r'<Relationship [^>]*calcChain[^>]*/>', '', files['xl/_rels/workbook.xml.rels'].decode()).encode()
with zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED) as z:
    for i in zin.infolist():
        if i.filename in files: z.writestr(i, files[i.filename])
print('từ điển: %d bảng, %d cột; quan hệ: %d; thiếu từ điển: %s' % (len(tables), len(dict_rows), len(rels), missing or 'không'))

# --- 5. Mermaid ---
if MMD:
    types = {(r[0], r[2]): (r[3] or 'string') for r in dict_rows}
    keys = {(r[0], r[2]): (r[4] or '') for r in dict_rows}
    card = {'1 : N': '||--o{', '1 : 0..1': '||--o|', '0..1 : N': '|o--o{'}
    L = ['erDiagram']
    for r in rels:
        L.append('  %s %s %s : "%s"' % (r[0], card.get(str(r[3]).split(' (')[0], '||--o{'), r[1], r[2]))
    for t in tables:
        L.append('  %s {' % t)
        for c in headers[t]:
            k = keys.get((t, c)); ty = types.get((t, c), 'string')
            L.append('    %s %s%s' % (ty, c.replace('-', '_'), (' ' + k) if k in ('PK', 'FK') else ''))
        L.append('  }')
    open(MMD, 'w').write('\n'.join(L) + '\n')
