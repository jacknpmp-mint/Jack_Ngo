"""Đọc Selling_Price_Database (R05+) → tools/schema_raw.json (cho gen_config.py) + test/data_real.json (lưới ô cho test mô phỏng).

Khác bộ khung chuẩn (extract_xlsx.py — tiêu đề dòng 1): file này có dòng 1 = tiêu đề sheet, dòng 2 = MÃ CỘT,
dòng 3 = nhãn (tiêu đề Excel Table), dữ liệu từ dòng 4; lược đồ lấy từ sheet DataDictionary; danh mục từ sheet Config
(dòng 7 = tên danh mục, dòng 8–37 = giá trị, cột A/C/E/…).

Dùng (từ thư mục app):  python3 -I tools/extract_spd.py <file.xlsx> [--rows 60]
"""
import json, os, re, sys, datetime, warnings, argparse
warnings.filterwarnings('ignore')
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
ap = argparse.ArgumentParser(); ap.add_argument('xlsx'); ap.add_argument('--rows', type=int, default=60)
a = ap.parse_args()
wf = openpyxl.load_workbook(a.xlsx)                  # công thức
wv = openpyxl.load_workbook(a.xlsx, data_only=True)  # giá trị lưu sẵn

TABLES = ['Staff', 'Markets', 'Customers', 'Items', 'Materials', 'FX_Rates', 'Material_Prices', 'Projects',
          'Project_Items', 'Project_Freight', 'Project_Inputs_Bid', 'Project_Inputs_Job']
LIST_BY_FIELD = {'currency': 'Currency', 'sell_currency': 'Currency', 'cost_currency': 'Currency',
                 'access_level': 'Access level', 'active': 'Active'}
STATIC_LISTS = {'Access level': ['Viewer', 'Editor', 'Admin'], 'Active': ['Y', 'N']}
TYPE = {'Text': 'string', 'List': 'string', 'Number': 'number', 'Percent': 'number', 'Date': 'date'}

# ---------------------------------------------------------------- danh mục Config
cfg = wv['Config']
lists = {}
for c in range(1, cfg.max_column + 1):
    name = cfg.cell(7, c).value
    if not name: continue
    name = re.sub(r'\s*\(.*\)$', '', str(name)).strip()           # 'Stage (row 1 = bid…)' → 'Stage'
    vals = [str(cfg.cell(r, c).value).strip() for r in range(8, 38) if cfg.cell(r, c).value not in (None, '')]
    lists[name] = vals
lists.update(STATIC_LISTS)
base_currency = cfg['B3'].value

# ---------------------------------------------------------------- DataDictionary
dd = wv['DataDictionary']
dic = {}
for r in range(4, dd.max_row + 1):
    row = [dd.cell(r, c).value for c in range(1, 10)]
    if not row[0] or not row[2]: continue
    dic[(row[0], row[2])] = dict(label=row[3], dtype=row[4], required=row[5] == 'Yes', key=row[6] or '', how=row[7] or '', desc=row[8] or '')

tables, warn = {}, []
for t in TABLES:
    ws = wf[t]
    codes = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    cols = []
    for i, code in enumerate(codes):
        if not code: continue
        d = dic.get((t, code))
        if not d: warn.append(f'{t}.{code}: không có trong DataDictionary'); d = dict(label=code, dtype='Text', required=False, key='', how='', desc='')
        fk = ''
        m = re.match(r'FK\s*>\s*(\w+)', d['key'])
        if m: fk = m.group(1)
        lst = ''
        m = re.search(r'Config\s*>\s*([A-Za-z ]+)', d['desc'])
        if m: lst = m.group(1).strip()
        if code in LIST_BY_FIELD and d['dtype'] == 'List': lst = LIST_BY_FIELD[code]
        if lst and lst not in lists: warn.append(f'{t}.{code}: danh mục "{lst}" không có trong Config'); lst = ''
        formula = d['how'] == 'Formula'
        f_rows = sum(1 for r in range(4, ws.max_row + 1) if isinstance(ws.cell(r, i + 1).value, str) and ws.cell(r, i + 1).value.startswith('='))
        if formula and f_rows <= 1: warn.append(f'{t}.{code}: DataDictionary ghi Formula nhưng chỉ {f_rows} dòng có công thức')
        cols.append(dict(name=code, type=TYPE.get(d['dtype'], 'string'), pct=d['dtype'] == 'Percent', fk=fk, list=lst,
                         formula=formula, formulaRows=f_rows if formula else 0, required=d['required'],
                         pk=d['key'] == 'PK', label=d['label'], desc=d['desc'], how=d['how']))  # desc: chỉ dùng làm ghi chú EN (make_spec_spd)
    pk = next((c['name'] for c in cols if c['pk']), cols[0]['name'])
    ids = [wv[t].cell(r, codes.index(pk) + 1).value for r in range(4, wv[t].max_row + 1)]
    ids = [x for x in ids if x not in (None, '')]
    tables[t] = dict(pk=pk, idPrefix='', idStyle='', rows=len(ids), columns=cols, sample=[str(x) for x in ids[:3]])

raw = dict(source=os.path.basename(a.xlsx), prefix='', headerRow=2, firstDataRow=4, baseCurrency=base_currency,
           tables=tables, lists=lists, warnings=warn)
json.dump(raw, open(os.path.join(HERE, 'schema_raw.json'), 'w'), ensure_ascii=False, indent=1)

# ---------------------------------------------------------------- lưới ô cho test (giữ công thức để app không ghi đè)
def cell(vf, vv):
    f = vf if isinstance(vf, str) and vf.startswith('=') else ''
    v = '' if f else vv
    if isinstance(v, datetime.datetime): v = {'$d': v.strftime('%Y-%m-%dT%H:%M:%S')}
    elif isinstance(v, datetime.date): v = {'$d': v.strftime('%Y-%m-%dT00:00:00')}
    if v is None: v = ''
    return {'v': v, 'f': f[1:] if f else ''} if f else {'v': v}

grid = {}
for t in TABLES + ['Config']:
    ws, wsv = wf[t], wv[t]
    last = min(ws.max_row, 3 + a.rows) if t != 'Config' else ws.max_row
    grid[t] = [[cell(ws.cell(r, c).value, wsv.cell(r, c).value) for c in range(1, ws.max_column + 1)] for r in range(1, last + 1)]
os.makedirs(os.path.join(ROOT, 'test'), exist_ok=True)
json.dump(grid, open(os.path.join(ROOT, 'test', 'data_real.json'), 'w'), ensure_ascii=False)

print(f'{raw["source"]}: {len(tables)} bảng, {sum(len(t["columns"]) for t in tables.values())} cột; base = {base_currency}')
for n, t in tables.items():
    fks = [c['name'] + '→' + c['fk'] for c in t['columns'] if c['fk']]
    print(f'  {n:20s} {len(t["columns"]):3d} cột · {t["rows"]:3d} dòng · PK={t["pk"]} {t["sample"]} · FK={", ".join(fks)}')
print('Danh mục:', {k: len(v) for k, v in lists.items()})
for w in warn: print('  ⚠ ' + w)
