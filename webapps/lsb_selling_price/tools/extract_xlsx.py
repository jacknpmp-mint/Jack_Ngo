"""Đọc file CSDL .xlsx bất kỳ (quy ước kiểu LSB) → tools/schema_raw.json + test/data_real.json.

Dùng (từ thư mục app, vd webapps/<app>/):
  python3 -I tools/extract_xlsx.py <file.xlsx> [--prefix HR_] [--out-data test/data_real.json]

Nhận diện (theo thứ tự ưu tiên, mọi suy đoán đều ghi vào "warnings" để hỏi lại người dùng):
  - Bảng: sheet có tên bắt đầu bằng tiền tố chung (tự tìm, hoặc --prefix). Bỏ qua Lists, ERD_Dictionary, Dashboard, Huong_dan…
  - Kiểu cột: ghi chú tiêu đề "Kiểu: string|number|date" → nếu không có thì đoán từ dữ liệu.
  - PK: tiêu đề màu cam A7650F → nếu không có: cột đầu tiên (cảnh báo).
  - FK: dropdown (x14 hoặc thường) trỏ tới <Bảng>!$A → nếu không có: tên cột trùng PK của bảng khác (cảnh báo, inferred).
  - Danh mục: dropdown trỏ tới Lists!$X (tên danh mục = tiêu đề cột X của sheet Lists).
  - AUTO: dòng 2 có công thức (ghi cả số dòng có công thức để phát hiện "công thức chỉ ở dòng 2").
  - Mẫu mã: giá trị PK hiện có (PJ-2026-001 → prefix PJ, idStyle year; ST-001 → catalog).
  - Ghi chú cột từ ERD_Dictionary (cột "Bảng", "Cột", "Ghi chú") nếu có.
Đọc trực tiếp XML của sheet vì openpyxl bỏ mất dropdown dạng x14.
"""
import json, re, sys, zipfile, datetime, warnings, os, argparse, collections
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
warnings.filterwarnings('ignore')

ap = argparse.ArgumentParser()
ap.add_argument('xlsx')
ap.add_argument('--prefix', default='')
ap.add_argument('--out-schema', default='')
ap.add_argument('--out-data', default='')
a = ap.parse_args()
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
OUT_SCHEMA = a.out_schema or os.path.join(HERE, 'schema_raw.json')
OUT_DATA = a.out_data or os.path.join(ROOT, 'test', 'data_real.json')

SKIP = {'lists', 'erd_dictionary', 'dashboard', 'huong_dan', 'hướng_dẫn', 'readme', 'guide', 'instructions'}
PK_FILL, FK_FILL, AUTO_FILL = 'A7650F', '1D5C96', '7A8894'

wb_f = openpyxl.load_workbook(a.xlsx)
wb_v = openpyxl.load_workbook(a.xlsx, data_only=True)
z = zipfile.ZipFile(a.xlsx)
wbxml = z.read('xl/workbook.xml').decode()
rels = z.read('xl/_rels/workbook.xml.rels').decode()
rid2t = {}
for m in re.finditer(r'<Relationship [^>]*>', rels):
    tag = m.group(0)
    i = re.search(r'Id="([^"]+)"', tag); t = re.search(r'Target="([^"]+)"', tag)
    if i and t: rid2t[i.group(1)] = t.group(1)
sheet_xml = {}
for m in re.finditer(r'<sheet [^>]*>', wbxml):
    tag = m.group(0)
    n = re.search(r'name="([^"]+)"', tag); r = re.search(r'r:id="([^"]+)"', tag)
    if n and r:
        target = rid2t[r.group(1)].lstrip('/')
        sheet_xml[n.group(1).replace('&amp;', '&')] = target if target.startswith('xl/') else 'xl/' + target

warn = []
names = [n for n in wb_f.sheetnames if n.lower() not in SKIP]
prefix = a.prefix
if not prefix:
    cnt = collections.Counter(n.split('_')[0] + '_' for n in names if '_' in n)
    prefix = cnt.most_common(1)[0][0] if cnt else ''
    if prefix: warn.append(f'Tiền tố tự nhận: {prefix} ({cnt[prefix]}/{len(names)} sheet). Sheet không theo tiền tố bị bỏ qua: ' +
                           ', '.join(n for n in names if not n.startswith(prefix)) if any(not n.startswith(prefix) for n in names) else f'Tiền tố tự nhận: {prefix}')
table_names = [n for n in names if n.startswith(prefix)] if prefix else names

lists, list_by_col = {}, {}
if 'Lists' in wb_f.sheetnames:
    ws = wb_f['Lists']
    for c in range(1, ws.max_column + 1):
        name = ws.cell(1, c).value
        if not name: continue
        list_by_col[get_column_letter(c)] = str(name).strip()
        vals = [v for v in (ws.cell(r, c).value for r in range(2, ws.max_row + 1)) if v not in (None, '')]
        lists[str(name).strip()] = list(dict.fromkeys(vals))

def cols_of(sqref):
    out = set()
    for part in sqref.split():
        m = re.match(r'\$?([A-Z]+)\$?\d*(?::\$?([A-Z]+)\$?\d*)?', part)
        if not m: continue
        x = column_index_from_string(m.group(1)); y = column_index_from_string(m.group(2) or m.group(1))
        out.update(range(x, y + 1))
    return out

def validations(xml):
    """[(formula, sqref)] cho cả dropdown x14 và dropdown thường."""
    out = []
    for m in re.finditer(r'<x14:dataValidation [^>]*type="list".*?<xm:f>([^<]*)</xm:f>.*?<xm:sqref>([^<]*)</xm:sqref>', xml, re.S):
        out.append((m.group(1), m.group(2)))
    for m in re.finditer(r'<dataValidation ([^>]*)>(.*?)</dataValidation>', xml, re.S):
        attrs, body = m.group(1), m.group(2)
        if 'type="list"' not in attrs: continue
        f = re.search(r'<formula1>([^<]*)</formula1>', body); sq = re.search(r'sqref="([^"]+)"', attrs)
        if f and sq: out.append((f.group(1).replace('&amp;', '&'), sq.group(1)))
    return out

# Ghi chú từ ERD_Dictionary
dic = {}
if 'ERD_Dictionary' in wb_f.sheetnames:
    ws = wb_f['ERD_Dictionary']
    head = [str(ws.cell(1, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
    def idx(*keys):
        for k in keys:
            if k in head: return head.index(k)
        return None
    it, ic, inote = idx('bảng', 'table'), idx('cột', 'column'), idx('ghi chú', 'note', 'notes')
    if it is not None and ic is not None:
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[it] and r[ic] and inote is not None and r[inote]:
                dic[(str(r[it]).strip(), str(r[ic]).strip())] = str(r[inote]).strip()

def guess_type(vals):
    vals = [v for v in vals if v not in (None, '')]
    if not vals: return 'string'
    if all(isinstance(v, (datetime.datetime, datetime.date)) for v in vals): return 'date'
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals): return 'number'
    return 'string'

ID_RE = re.compile(r'^([A-Z]{1,5})-(?:(\d{4})-)?(\d{2,5})$')
tables, data = {}, {}
for name in table_names:
    ws, wv = wb_f[name], wb_v[name]
    xml = z.read(sheet_xml[name]).decode()
    fk, lst = {}, {}
    for f, sq in validations(xml):
        mm = re.match(r"=?'?([^'!]+)'?!\$A\$", f)
        ml = re.match(r"=?'?Lists'?!\$([A-Z]+)\$", f)
        for c in cols_of(sq):
            if ml: lst[c] = list_by_col.get(ml.group(1), '')
            elif mm and mm.group(1) in table_names: fk[c] = mm.group(1)
    cols = []
    last = 1
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(r, c).value not in (None, '') and not (isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith('='))
               for c in range(1, ws.max_column + 1)):
            last = r; break
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h in (None, ''): continue
        cell = ws.cell(1, c)
        note = cell.comment.text if cell.comment else ''
        mt = re.search(r'Kiểu:\s*(\w+)', note)
        fill = ''
        try: fill = (cell.fill.fgColor.rgb or '')[-6:].upper() if cell.fill and cell.fill.fgColor and isinstance(cell.fill.fgColor.rgb, str) else ''
        except Exception: fill = ''
        f_rows = [r for r in range(2, ws.max_row + 1) if isinstance(ws.cell(r, c).value, str) and str(ws.cell(r, c).value).startswith('=')]
        vals = [wv.cell(r, c).value for r in range(2, last + 1)]
        cols.append(dict(name=str(h).strip(), col=get_column_letter(c), type=mt.group(1) if mt else guess_type(vals), typeFromNote=bool(mt),
                         fill=fill, fk=fk.get(c, ''), list=lst.get(c, ''), formula=2 in f_rows, formulaRows=len(f_rows),
                         note=dic.get((name, str(h).strip()), ''), headerNote=note.strip()))
    # PK
    pk = [x for x in cols if x['fill'] == PK_FILL]
    if not pk and cols:
        warn.append(f'{name}: không có tiêu đề màu PK (cam {PK_FILL}) → tạm coi cột đầu "{cols[0]["name"]}" là PK.')
        cols[0]['pkInferred'] = True
    elif len(pk) > 1:
        warn.append(f'{name}: nhiều cột màu PK ({", ".join(x["name"] for x in pk)}) → dùng cột đầu.')
    pkname = (pk[0] if pk else cols[0])['name'] if cols else ''
    # mẫu mã
    ids = [wv.cell(r, [x['name'] for x in cols].index(pkname) + 1).value for r in range(2, ws.max_row + 1)] if cols else []
    ids = [str(v) for v in ids if v not in (None, '')]
    pat = collections.Counter()
    for v in ids:
        m = ID_RE.match(v)
        if m: pat[(m.group(1), 'year' if m.group(2) else 'catalog')] += 1
    sample = pat.most_common(1)[0][0] if pat else None
    for x in cols:
        if x['formula'] and x['formulaRows'] <= 1:
            warn.append(f'{name}.{x["name"]}: công thức CHỈ ở dòng 2 ({x["formulaRows"]} dòng) — không coi là AUTO; cân nhắc calc (app tự tính khi để trống) hoặc kéo công thức xuống.')
    tables[name] = dict(pk=pkname, idPrefix=sample[0] if sample else '', idStyle=sample[1] if sample else '', rows=len(ids), columns=cols)
    grid = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            fv = ws.cell(r, c).value; vv = wv.cell(r, c).value
            f = fv if isinstance(fv, str) and fv.startswith('=') else ''
            v = vv if f else fv
            if isinstance(v, (datetime.datetime, datetime.date)):
                if not isinstance(v, datetime.datetime): v = datetime.datetime(v.year, v.month, v.day)
                v = {'$d': v.strftime('%Y-%m-%dT%H:%M:%S')}
            if isinstance(v, datetime.time): v = v.strftime('%H:%M:%S')
            row.append({'v': v, 'f': f})
        grid.append(row)
    data[name] = grid

# FK suy ra từ tên cột trùng PK bảng khác
pks = {t['pk']: n for n, t in tables.items()}
for n, t in tables.items():
    for x in t['columns']:
        if x['name'] == t['pk'] or x['fk']: continue
        if x['name'] in pks and pks[x['name']] != n:
            x['fk'] = pks[x['name']]; x['fkInferred'] = True
            warn.append(f'{n}.{x["name"]}: không có dropdown FK → suy ra FK tới {x["fk"]} (trùng tên PK).')
    for x in t['columns']:
        if x['fill'] == FK_FILL and not x['fk']:
            warn.append(f'{n}.{x["name"]}: tiêu đề màu FK nhưng không tìm được bảng cha → cần chỉ định trong app_spec.json (tables.{n}.fk).')
    if not t['idPrefix'] and t['rows']:
        warn.append(f'{n}: mã PK hiện có không theo mẫu XX-NĂM-STT / XX-STT → cần chỉ định prefix + idStyle.')

json.dump({'source': os.path.basename(a.xlsx), 'prefix': prefix, 'tables': tables, 'lists': lists, 'warnings': warn,
           'otherSheets': [n for n in wb_f.sheetnames if n not in tables]},
          open(OUT_SCHEMA, 'w'), ensure_ascii=False, indent=1)
os.makedirs(os.path.dirname(OUT_DATA), exist_ok=True)
json.dump(data, open(OUT_DATA, 'w'), ensure_ascii=False)
print(f'tiền tố {prefix or "(không)"} · bảng: {len(tables)} · cột: {sum(len(t["columns"]) for t in tables.values())} · danh mục: {len(lists)}')
for n, t in tables.items():
    print(f'  {n:34s} {len(t["columns"]):3d} cột · {t["rows"]:3d} dòng · PK={t["pk"]} ({t["idPrefix"] or "?"}/{t["idStyle"] or "?"}) · FK=' +
          ','.join(x['name'] + '→' + x['fk'] + ('?' if x.get('fkInferred') else '') for x in t['columns'] if x['fk']))
if warn:
    print('CẦN XÁC NHẬN:')
    for w in warn: print('  ⚠ ' + w)
