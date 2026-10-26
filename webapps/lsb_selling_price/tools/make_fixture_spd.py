"""Dữ liệu MINH HOẠ cho test (test/data.json): cấu trúc đúng file R05 (dòng 1–3, công thức giữ nguyên),
giá trị nhập tay giả lập (khách hàng / dự án / giá không phải số thật). Sinh cố định (seed) để test lặp lại được.

Dùng (từ thư mục app):  python3 -I tools/make_fixture_spd.py <R05.xlsx>
"""
import json, os, sys, random, datetime, warnings
warnings.filterwarnings('ignore')
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
wf = openpyxl.load_workbook(sys.argv[1])
wv = openpyxl.load_workbook(sys.argv[1], data_only=True)
rnd = random.Random(7)
D = lambda s: datetime.datetime.strptime(s, '%Y-%m-%d')

customers = [('CUS-0001', 'Khach hang A (minh hoa)', 'General contractor', 'Vietnam'), ('CUS-0002', 'Khach hang B (minh hoa)', 'Owner / Investor', 'Indonesia'),
             ('CUS-0003', 'Khach hang C (minh hoa)', 'Dealer', 'Thailand'), ('CUS-0004', 'Khach hang D (minh hoa)', 'General contractor', 'Philippines'),
             ('CUS-0005', 'Khach hang E (minh hoa)', 'Owner / Investor', 'Angola'), ('CUS-0006', 'Khach hang F (minh hoa)', 'Dealer', 'Australia'),
             ('CUS-0007', 'Khach hang G (minh hoa)', 'General contractor', 'Vietnam')]
items = [('ITM-01', 'Primary framing', 'kg', 'Steel plate'), ('ITM-02', 'Secondary framing - purlins', 'kg', 'Cold-formed steel - purlins'),
         ('ITM-03', 'Roof sheeting', 'm2', 'Steel sheeting'), ('ITM-04', 'Wall sheeting', 'm2', 'Steel sheeting'),
         ('ITM-05', 'Anchor bolts', 'kg', 'Anchor bolts'), ('ITM-06', 'Rod bracing', 'kg', 'Rod Bracing')]
materials = [('MAT-0001', 'Hot-rolled steel plate', 'Steel plate', 'kg', 15800), ('MAT-0002', 'Galvanised purlin coil', 'Cold-formed steel - purlins', 'kg', 19500),
             ('MAT-0003', 'Pre-painted steel sheet 0.45', 'Steel sheeting', 'm2', 98000), ('MAT-0004', 'Anchor bolt M24', 'Anchor bolts', 'kg', 31000),
             ('MAT-0005', 'Round bar bracing', 'Rod Bracing', 'kg', 17200)]
months = ['%d-%02d-01' % (y, m) for y in (2025, 2026) for m in range(1, 13) if not (y == 2026 and m > 10)]
fx = []
for i, d in enumerate(months):
    fx.append((d, 'USD', round(25300 + i * 45 + rnd.random() * 60)))
    fx.append((d, 'IDR', round(1.55 + rnd.random() * 0.08, 3)))
    fx.append((d, 'PHP', round(435 + i * 1.2 + rnd.random() * 8)))
fx = [('FX-%04d' % (k + 1),) + x for k, x in enumerate(fx)]

def rate(cur, date):
    if cur == 'VND': return 1
    r = sorted([x for x in fx if x[2] == cur and x[1] <= date], key=lambda x: x[1])
    return r[-1][3] if r else None

prices = []
for m in materials:
    for k, sp in enumerate(['Spec A', 'Spec B']):
        for mi in (0, 8, 16):
            d = months[mi + (2 if k else 0)]
            prices.append(('PRC-%05d' % (len(prices) + 1), m[0], m[1] + ' - ' + sp, d, 'NCC minh hoa', 'VND', round(m[4] * (1 + k * 0.06) * (1 + mi * 0.004) * (0.97 + rnd.random() * 0.06))))
names = ['Nha xuong Binh Duong', 'Kho lanh Long An', 'Nha may Cikarang', 'Warehouse Rayong', 'Factory Laguna', 'Kho Luanda', 'Shed Brisbane',
         'Nha xuong Bac Ninh', 'Plant Batam', 'Hangar Cebu', 'Kho Dong Nai', 'Factory Chonburi']
mk = ['TT-01', 'TT-01', 'TT-02', 'TT-03', 'TT-04', 'TT-05', 'TT-06', 'TT-01', 'TT-02', 'TT-04', 'TT-01', 'TT-03']
sts = ['Completed', 'Won', 'Lost', 'Won', 'Quoting', 'Lost', 'Won', 'Completed', 'Quoting', 'Won', 'Lost', 'Quoting']
cur_by = {'TT-01': 'VND', 'TT-02': 'USD', 'TT-03': 'USD', 'TT-04': 'PHP', 'TT-05': 'USD', 'TT-06': 'USD'}
projects, lines, freight, inb, inj = [], [], [], [], []
for i, n in enumerate(names):
    pid = 'PRJ-%04d' % (i + 1)
    qd = months[min(len(months) - 1, int(i * 1.8) + 1)][:8] + '%02d' % (5 + rnd.randrange(20))
    won = sts[i] in ('Won', 'Completed')
    sc = 'IDR' if i == 8 else cur_by[mk[i]]
    fd = '2024-12-20' if i == 5 else qd                      # PRJ-0006: không có tỷ giá USD ≤ ngày → cảnh báo
    cd = (qd[:5] + '%02d' % min(12, int(qd[5:7]) + 1) + '-15') if won else ''
    projects.append([pid, n + ' (minh hoa)', customers[i % 7][0], mk[i], qd, cd, sts[i], ('J26-%d' % (i + 10)) if won and i != 9 else '', sc, 'VND', fd, 'Estimator A'])
    if i == 11: continue                                      # PRJ-0012: chưa có dòng item Bid → cảnh báo
    for it in [x for k, x in enumerate(items) if k < 4 or rnd.random() > 0.4]:
        mat = next((m for m in materials if m[2] == it[3]), materials[0])
        b = {'kg': [mat[4], 3200, 2600, 600], 'm2': [mat[4], 9000, 14000, 2500]}[it[2]]
        cm, cf, ce, co = round(b[0] * (0.94 + rnd.random() * 0.14)), round(b[1] * (0.9 + rnd.random() * 0.2)), round(b[2] * (0.9 + rnd.random() * 0.2)), b[3]
        gm = 0.12 + rnd.random() * 0.14
        r = rate(sc, fd) or 1
        sell = round((cm + cf + ce + co) / (1 - gm) / r, 0 if sc == 'VND' else 3)
        lines.append(['PIT-%05d' % (len(lines) + 1), pid, it[0], 'Bid', cm, cf, ce, co, sell])
        if won and i != 9:
            d = 0.97 + rnd.random() * 0.09
            lines.append(['PIT-%05d' % (len(lines) + 1), pid, it[0], 'Job', round(cm * d), round(cf * (0.97 + rnd.random() * 0.08)), ce, co, sell])
    if i != 6:
        freight.append(['FRT-%05d' % (len(freight) + 1), pid, 'DAP' if mk[i] == 'TT-01' else rnd.choice(['FOB', 'CFR', 'CIF']), 'Cang dich (minh hoa)', "40' container",
                        'VND' if mk[i] == 'TT-01' else 'USD', round(9e6 + rnd.random() * 6e6) if mk[i] == 'TT-01' else round(1400 + rnd.random() * 2600)])
    for k, m in enumerate(materials):
        if k > 2 and rnd.random() < 0.5: continue
        pr = sorted([p for p in prices if p[1] == m[0] and p[3] <= fd and 'Spec A' in p[2]], key=lambda p: p[3])
        pr = pr[-1] if pr else None
        sp = pr[2] if pr else m[1] + ' - Spec A'
        bp = round((pr[6] if pr else m[4]) * (0.98 + rnd.random() * 0.06))
        inb.append(['INB-%05d' % (len(inb) + 1), pid, m[0], sp, pr[0] if pr else '', 'VND', bp])
        if won and i != 9:
            spj = sp + ' rev.1' if (k == 2 and i == 3) else sp             # PRJ-0004: lệch spec → Kiểm tra dữ liệu
            inj.append(['INJ-%05d' % (len(inj) + 1), pid, m[0], spj, cd, pr[0] if pr else '', 'VND', round(bp * (0.96 + rnd.random() * 0.12))])

FIELDS = {
    'Customers': (['customer_id', 'customer_name', 'customer_type', 'country'], customers),
    'Items': (['item_id', 'item_name', 'unit', 'material_group'], items),
    'Materials': (['material_id', 'material_name', 'material_group', 'unit'], [m[:4] for m in materials]),
    'FX_Rates': (['fx_id', 'effective_date', 'currency', 'rate'], fx),
    'Material_Prices': (['price_id', 'material_id', 'specification', 'effective_date', 'supplier', 'currency', 'unit_price'], prices),
    'Projects': (['project_id', 'project_name', 'customer_id', 'market_id', 'quote_date', 'contract_date', 'status', 'job_no', 'sell_currency', 'cost_currency', 'fx_date', 'owner'], projects),
    'Project_Items': (['line_id', 'project_id', 'item_id', 'stage', 'cost_material', 'cost_fabrication', 'cost_erection', 'cost_other', 'unit_sell'], lines),
    'Project_Freight': (['freight_id', 'project_id', 'incoterm', 'destination_port', 'cost_basis', 'currency', 'freight_cost'], freight),
    'Project_Inputs_Bid': (['input_id', 'project_id', 'material_id', 'specification', 'price_id', 'currency', 'unit_price'], inb),
    'Project_Inputs_Job': (['input_id', 'project_id', 'material_id', 'specification', 'price_date', 'price_id', 'currency', 'unit_price'], inj),
}
DATES = {'effective_date', 'quote_date', 'contract_date', 'fx_date', 'price_date'}

def cell(vf, vv):
    if isinstance(vf, str) and vf.startswith('='): return {'v': '', 'f': vf[1:]}
    if isinstance(vv, datetime.datetime): return {'v': {'$d': vv.strftime('%Y-%m-%dT%H:%M:%S')}}
    return {'v': '' if vv is None else vv}

grid = {}
for t in ['Staff', 'Markets', 'Customers', 'Items', 'Materials', 'FX_Rates', 'Material_Prices', 'Projects', 'Project_Items',
          'Project_Freight', 'Project_Inputs_Bid', 'Project_Inputs_Job', 'Config']:
    ws, wsv = wf[t], wv[t]
    n = len(FIELDS.get(t, ([], []))[1])
    last = ws.max_row if t == 'Config' else min(ws.max_row, 3 + max(n, 1) + 25)
    g = [[cell(ws.cell(r, c).value, wsv.cell(r, c).value) for c in range(1, ws.max_column + 1)] for r in range(1, last + 1)]
    if t in FIELDS:
        codes = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        cols, rows = FIELDS[t]
        for k, row in enumerate(rows):
            for name, v in zip(cols, row):
                if v == '' or v is None: continue
                if name in DATES: v = {'$d': v + 'T00:00:00'}
                g[3 + k][codes.index(name)] = {'v': v}
    grid[t] = g
json.dump(grid, open(os.path.join(ROOT, 'test', 'data.json'), 'w'), ensure_ascii=False)
print('test/data.json:', {k: len(v[1]) for k, v in FIELDS.items()})
