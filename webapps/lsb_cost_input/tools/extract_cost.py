"""Đọc Cost_Input_Database_Part2 (R03) → tools/schema_raw.json + test/data_real.json (cùng định dạng extract_xlsx.py).

Khác quy ước LSB nên dùng script riêng thay cho extract_xlsx.py:
  - Mỗi sheet: dòng 1 tiêu đề, dòng 2 = MÃ CỘT (app đọc theo dòng này), dòng 3 = nhãn hiển thị, dữ liệu từ dòng 4.
  - Lược đồ lấy từ sheet DataDictionary (kiểu, bắt buộc, PK/FK, "Formula" = cột AUTO, mô tả = ghi chú).
  - Danh sách chọn ở sheet Config: dòng 7 = tên danh sách, dòng 8–107 = giá trị (cột A, C, E…); dropdown trỏ Config!$X$8:$X$107.
  - Sheet Staff (R03) chưa có trong DataDictionary của R02 → lấy từ DataDictionary R03 (đã thêm dòng).
Dùng (từ thư mục app):  python3 -I tools/extract_cost.py <R03.xlsx> [--rows 40]
  --rows: số dòng mỗi sheet giữ trong test/data_real.json (dòng trống có công thức vẫn giữ để mô phỏng dòng dựng sẵn).
"""
import json, os, re, sys, zipfile, datetime, argparse, warnings
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
warnings.filterwarnings('ignore')

ap = argparse.ArgumentParser()
ap.add_argument('xlsx'); ap.add_argument('--rows', type=int, default=40)
a = ap.parse_args()
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
HEADER_ROW, FIRST_ROW = 2, 4
SKIP = {'README', 'Config', 'DataDictionary'}

wf = openpyxl.load_workbook(a.xlsx)
wv = openpyxl.load_workbook(a.xlsx, data_only=True)

# ----------------------------------------------------------- danh sách chọn (Config)
cfg = wf['Config']
lists, list_by_col = {}, {}
for c in range(1, cfg.max_column + 1):
    name = cfg.cell(7, c).value
    if not name: continue
    name = str(name).strip()
    list_by_col[get_column_letter(c)] = name
    lists[name] = list(dict.fromkeys(v for v in (cfg.cell(r, c).value for r in range(8, 108)) if v not in (None, '')))
lists['Access level'] = ['Viewer', 'Editor', 'Admin']
lists['Active'] = ['Y', 'N']

# ----------------------------------------------------------- DataDictionary
dd = wf['DataDictionary']
head = [str(x.value or "").strip().lower() for x in dd[3]]
def ix(k): return next(i for i, h in enumerate(head) if h.startswith(k))
I = dict(t=ix('table'), f=ix('field code'), l=ix('field label'), ty=ix('data type'), rq=ix('required'), k=ix('key'), how=ix('how'), d=ix('description'))
dic = {}
for r in dd.iter_rows(min_row=4, values_only=True):
    if not r[I['t']] or not r[I['f']]: continue
    dic.setdefault(str(r[I['t']]).strip(), {})[str(r[I['f']]).strip()] = r

TYPE = {'text': 'string', 'list': 'string', 'number': 'number', 'percent': 'number', 'date': 'date'}
warn = []
tables, data = {}, {}
for name in wf.sheetnames:
    if name in SKIP: continue
    ws, vs = wf[name], wv[name]
    if name not in dic: warn.append(f'{name}: không có trong DataDictionary → bỏ qua'); continue
    # dropdown → danh sách / FK
    lst, fk_dv = {}, {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != 'list' or not dv.formula1: continue
        f = str(dv.formula1)
        m = re.match(r"=?'?Config'?!\$([A-Z]+)\$8", f)
        mt = re.match(r"=?'?([A-Za-z_]+)'?!\$A\$4", f)
        for rng in str(dv.sqref).split():
            m2 = re.match(r'([A-Z]+)\d*(?::([A-Z]+)\d*)?', rng)
            for ci in range(column_index_from_string(m2.group(1)), column_index_from_string(m2.group(2) or m2.group(1)) + 1):
                if m: lst[ci] = list_by_col.get(m.group(1), '')
                elif mt: fk_dv[ci] = mt.group(1)
                elif name == 'Staff': lst[ci] = 'Access level' if get_column_letter(ci) == 'D' else 'Active'
    last = HEADER_ROW
    for r in range(ws.max_row, FIRST_ROW - 1, -1):
        if ws.cell(r, 1).value not in (None, ''): last = r; break
    cols, pk = [], ''
    for c in range(1, ws.max_column + 1):
        code = ws.cell(HEADER_ROW, c).value
        if code in (None, ''): continue
        code = str(code).strip()
        d = dic[name].get(code)
        if not d: warn.append(f'{name}.{code}: không có trong DataDictionary'); continue
        key = str(d[I['k']] or '')
        is_formula = str(d[I['how']] or '').lower() == 'formula'
        f_rows = sum(1 for r in range(FIRST_ROW, ws.max_row + 1) if isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith('='))
        fk = ''
        mfk = re.match(r'FK\s*>\s*(\w+)', key)
        if mfk: fk = mfk.group(1)
        elif c in fk_dv: fk = fk_dv[c]
        if key.startswith('PK') and not pk and (name != 'FX_Rates' or code == 'fx_id'): pk = code
        cols.append(dict(name=code, col=get_column_letter(c), type=TYPE.get(str(d[I['ty']]).strip().lower(), 'string'), typeFromNote=True,
                         fill='', fk=fk, list=lst.get(c, ''), formula=is_formula, formulaRows=f_rows if is_formula else 0,
                         note=str(d[I['d']] or '').strip(), headerNote='', label=str(d[I['l']] or '').strip(),
                         required=str(d[I['rq']] or '').strip().lower() == 'yes', pct=str(d[I['ty']]).strip().lower() == 'percent'))
        if is_formula and f_rows < 2: warn.append(f'{name}.{code}: DataDictionary ghi Formula nhưng sheet chỉ có {f_rows} dòng công thức')
    if not pk: warn.append(f'{name}: không tìm thấy PK'); pk = cols[0]['name']
    ids = [str(ws.cell(r, [x['name'] for x in cols].index(pk) + 1 + 0).value) for r in range(FIRST_ROW, last + 1)
           if ws.cell(r, [x['name'] for x in cols].index(pk) + 1).value not in (None, '')]
    pat = [re.match(r'^([A-Z]+)-(\d+)$', x) for x in ids]
    pat = [m for m in pat if m]
    tables[name] = dict(pk=pk, idPrefix=pat[0].group(1) if pat else '', idStyle='catalog', idDigits=len(pat[0].group(2)) if pat else 0,
                        rows=len(ids), columns=cols)
    # lưới dữ liệu cho mô phỏng: dòng 1..max(rows, last+5)
    nrows = min(ws.max_row, max(a.rows, last + 5))
    grid = []
    for r in range(1, nrows + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            fv = ws.cell(r, c).value; vv = vs.cell(r, c).value
            f = fv if isinstance(fv, str) and fv.startswith('=') else ''
            v = vv if f else fv
            if f and v is None: v = ''
            if isinstance(v, (datetime.datetime, datetime.date)):
                if not isinstance(v, datetime.datetime): v = datetime.datetime(v.year, v.month, v.day)
                v = {'$d': v.strftime('%Y-%m-%dT%H:%M:%S')}
            if isinstance(v, datetime.time): v = v.strftime('%H:%M:%S')
            row.append({'v': v, 'f': f})
        grid.append(row)
    data[name] = grid

# Config: giữ nguyên lưới (app đọc danh sách trực tiếp)
cg = []
for r in range(1, 108):
    cg.append([{'v': cfg.cell(r, c).value if cfg.cell(r, c).value is not None else '', 'f': ''} for c in range(1, cfg.max_column + 1)])
data['Config'] = cg

json.dump({'source': os.path.basename(a.xlsx), 'prefix': '', 'tables': tables, 'lists': lists, 'warnings': warn,
           'otherSheets': [n for n in wf.sheetnames if n not in tables], 'headerRow': HEADER_ROW, 'firstRow': FIRST_ROW,
           'baseCurrency': cfg['B3'].value},
          open(os.path.join(HERE, 'schema_raw.json'), 'w'), ensure_ascii=False, indent=1)
os.makedirs(os.path.join(ROOT, 'test'), exist_ok=True)
json.dump(data, open(os.path.join(ROOT, 'test', 'data_real.json'), 'w'), ensure_ascii=False)
print(f'bảng: {len(tables)} · cột: {sum(len(t["columns"]) for t in tables.values())} · danh sách: {len(lists)} · tiền gốc: {cfg["B3"].value}')
for n, t in tables.items():
    print(f'  {n:22s} {len(t["columns"]):3d} cột · {t["rows"]:2d} dòng · PK={t["pk"]} ({t["idPrefix"]}-{"0"*t["idDigits"]}) · FK=' +
          ','.join(x['name'] + '→' + x['fk'] for x in t['columns'] if x['fk']) + ' · DS=' + ','.join(sorted({x['list'] for x in t['columns'] if x['list']})))
for w in warn: print('  ⚠ ' + w)
